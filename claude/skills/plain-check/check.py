#!/usr/bin/env python3
"""Mechanical half of the plain-register check.

Finds the things an eye skips: exact sentence lengths, punctuation pivots,
paraprosdokian shapes, undefined-term candidates, dangling references, and
missing verbs. It reports candidates. Judging them is the model's job, because none of these tests can
tell a real problem from a false positive on its own.

Labels are separated from prose, because the register rules exempt headings,
card titles, chart captions and footers. Sentence-level checks run on prose
only. Term checks run on everything, because a label can carry jargon too.

usage:  check.py <file> [<file> ...]
        check.py --stdin          read text from stdin
        check.py --verbose        show every candidate, not the first 40
Reads .md, .txt, .docx, .pptx.
"""
import os, re, sys, zipfile
from collections import Counter

# ---------------------------------------------------------------- extraction
# Each extractor yields (kind, text) where kind is 'prose' or 'label'.

LABEL_MAX_WORDS = 8


def looks_like_label(s):
    """Short, no terminal punctuation, or all-caps styling."""
    s = s.strip()
    if not s:
        return True
    if len(s.split()) <= LABEL_MAX_WORDS and not re.search(r'[.!?]\s*$', s):
        return True
    letters = re.sub(r'[^A-Za-z]', '', s)
    return bool(letters) and letters.isupper() and len(s.split()) <= 12


def from_docx(p):
    with zipfile.ZipFile(p) as z:
        xml = z.read('word/document.xml').decode('utf8', 'ignore')
    heading = re.compile(r'w:pStyle w:val="(Heading|Title|Subtitle)')
    for para in re.split(r'</w:p>', xml):
        txt = re.sub(r'<[^>]+>', '', re.sub(r'<w:tab[^>]*/>', ' ', para)).strip()
        if txt:
            yield ('label' if heading.search(para) or looks_like_label(txt) else 'prose', txt)


def from_pptx(p):
    from pptx import Presentation
    for s in Presentation(p).slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                is_title = False
                try:
                    is_title = sh.is_placeholder and 'TITLE' in str(sh.placeholder_format.type)
                except Exception:
                    pass
                for line in sh.text_frame.text.split('\n'):
                    if line.strip():
                        yield ('label' if is_title or looks_like_label(line) else 'prose',
                               line.strip())
            if getattr(sh, 'has_table', False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        if c.text.strip():
                            yield ('label', c.text.strip())
        if s.has_notes_slide and s.notes_slide.notes_text_frame:
            for line in s.notes_slide.notes_text_frame.text.split('\n'):
                if line.strip():
                    yield ('label' if looks_like_label(line) else 'prose', line.strip())


def from_text(p):
    t = open(p, errors='ignore').read()
    if t.startswith('---'):                                   # drop frontmatter
        parts = t.split('---', 2)
        if len(parts) > 2:
            t = parts[2]
    t = re.sub(r'```.*?```', ' ', t, flags=re.S)               # fenced code
    t = re.sub(r'`[^`]+`', ' ', t)                             # inline code
    t = re.sub(r'!?\[([^\]]*)\]\([^)]*\)', r'\1', t)           # links, keep text
    for line in t.split('\n'):
        line = line.strip()
        if not line or re.match(r'^[-=|>*_\s]+$', line):
            continue
        if line.startswith('|'):                               # table row
            yield ('label', re.sub(r'\|', ' ', line).strip())
            continue
        if re.match(r'^#{1,6}\s', line):
            yield ('label', re.sub(r'^#{1,6}\s+', '', line))
            continue
        body = re.sub(r'^([-*+]|\d+\.)\s+', '', line)
        yield ('label' if looks_like_label(body) else 'prose', body)


def extract(p):
    ext = os.path.splitext(p)[1].lower()
    if ext == '.docx':
        return list(from_docx(p))
    if ext == '.pptx':
        return list(from_pptx(p))
    return list(from_text(p))


def drop_repeats(chunks):
    """Footers and running headers repeat. Count them once."""
    seen = Counter(t for _, t in chunks)
    out, used = [], set()
    for kind, t in chunks:
        if seen[t] >= 3:
            if t in used:
                continue
            used.add(t)
            kind = 'label'
        out.append((kind, t))
    return out

# ------------------------------------------------------------------ measures

VOWELS = 'aeiouy'


def syl(w):
    w = re.sub(r'[^a-z]', '', w.lower())
    if not w:
        return 0
    n, prev = 0, False
    for c in w:
        cur = c in VOWELS
        if cur and not prev:
            n += 1
        prev = cur
    if w.endswith('e') and n > 1:
        n -= 1
    return max(n, 1)


def split_sentences(text):
    return [s.strip() for s in re.split(r'(?<=[.!?])\s+', text) if s.strip()]

# ------------------------------------------------------------------- lexicons

# Irregular and very common verbs, in the forms that appear in prose. The
# -ed / -ing / -s tests below catch regular verbs, so this list only has to
# cover what those tests miss.
VERBS = set("""am is are was were be been being have has had do does did done
can could will would shall should may might must ought need needs dare
go goes went gone come comes came become becomes became get gets got gotten
make makes made take takes took taken give gives gave given put puts
say says said tell tells told see sees saw seen know knows knew known
think thinks thought find finds found feel feels felt leave leaves left
bring brings brought buy buys bought build builds built catch catches caught
choose chooses chose chosen cut cuts draw draws drew drawn drive drives drove
eat eats ate fall falls fell fight fights fought fly flies flew forget forgets
forgot grow grows grew hold holds held keep keeps kept lead leads led
lose loses lost mean means meant meet meets met pay pays paid read reads
ride rides rode ring rings rang rise rises rose run runs ran sell sells sold
send sends sent set sets shake shakes shook shine shines shone shoot shoots
show shows shown shut shuts sing sings sang sit sits sat sleep sleeps slept
speak speaks spoke spend spends spent stand stands stood strike strikes struck
swim swims swam teach teaches taught tear tears tore throw throws threw
understand understands understood wear wears wore win wins won write writes
wrote written let lets hurt hurts cost costs hit hits quit quits split splits
spread spreads seem seems appear appears remain remains stay stays turn turns
look looks want wants like likes use uses work works help helps start starts
call calls try tries ask asks move moves live lives believe believes bring
happen happens provide provides sit stand lose add adds change changes
follow follows create creates open opens walk walks offer offers remember
consider considers expect expects allow allows serve serves die dies send
build fill fills reach reaches kill kills raise raises pass passes decide
return returns explain explains hope hopes develop develops carry carries
break breaks receive receives agree agrees support supports hit produce
produces eat cover covers catch draw choose cause causes point points
listen listens realize realizes place places close closes involve involves
increase increases reduce reduces name names replace replaces enter enters
share shares apply applies check checks report reports treat treats
account accounts matter matters count counts scale scales ship ships
own owns run sit hold price prices model models map maps flag flags
land lands log logs test tests track tracks fund funds staff staffs""".split())

SUFFIX_VERB = re.compile(r'\w{3,}(ed|ing)$')

# Long words common enough not to need defining.
COMMON_LONG = set("""another available business company customer development
different difficult document education example experience following generally
government important including information interest management necessary
operation organization particular performance possible probably production
question relationship remember several similar something sometimes technology
together understand usually already another anything beautiful community
computer consider continue conversation decision delivery difference direction
discussion economy elsewhere employee especially everybody everyone everything
family financial however industry material mechanism national natural
opportunity otherwise personal physical political popular position positive
potential practical presentation president property provided quality quickly
recently regular remaining resource responsible security services situation
specific standard strategy suddenly suggested supported terrible together
transformation typically university visitor whatever whenever wherever
yesterday company companies percent million billion approved agreement
digital digitize digitized computer computers actually prototype prototypes
customer customers company companies process processes system systems
document documents record records contract contracts service services
version versions permission permissions comment comments customer employee
manager managers business businesses product products project projects
category categories activity activities capital capacity clarity quality
identity ability policy policies summary category deliver delivery
another company anything everyone everything already however therefore
simply usually finally recently currently previously obviously certainly
enterprise enterprises internal external original financial personal
critical practical technical physical logical general generally
industry industries agency agencies country countries family families
history histories inventory inventories memory memories energy energies
example examples estimate estimates evidence experience experiences
however whatever whichever incident incidents interest interests
manager operator operators partner partners provider providers
regular regularly relative relatively separate separately similar
attention condition conditions decision decisions direction directions
discussion education function functions location locations option options
portion position positions question questions reaction relation relations
section sections solution solutions station stations version
computer telephone telephones together tomorrow yesterday
""".split())

# Nouns I reach for that stand in for a real thing instead of naming it.
METAPHOR_NOUNS = """machinery envelope object substrate spine lens rail beat
wedge dial fabric plumbing scaffolding scaffold engine arc thread lever ceiling
runway moat flywheel anchor backbone funnel seam texture grain footprint
posture appetite muscle bandwidth north_star centre_of_gravity long_pole
surface_area""".split()

# Terms of art built from ordinary short words, so no length test can see them.
JARGON_PHRASES = """data room|arrangement agreement|cap rate|run rate|term sheet|
reverse fee|termination fee|break fee|earn out|drag along|tag along|due diligence|
fairness opinion|going concern|net new|land and expand|book of business|rate card|
change order|scope creep|statement of work|steady state|end state|target state|
current state|source of truth|system of record|golden record|data lake|data mesh|
feature store|context window|token budget|human in the loop|guard rail|red team|
ground truth|hallucination|fine tune|prompt injection|digital twin|control tower|
value pool|no regret|quick win|north star|operating model|capability map|
maturity model|cost to serve|total cost of ownership|build versus buy|
capex|opex|run book|air gap|single pane of glass|left shift|shift left""".replace('\n', '').split('|')

LEADING_REF = re.compile(r'^(It|This|That|These|Those|They|Such|Here|There)\b')
SAME_X = re.compile(r'\bthe (same|latter|former|above|underlying|resulting) \w+', re.I)


FRAGMENT_START = re.compile(
    r'^(The|A|An|From|With|In|On|At|For|After|Before|Under|Behind|Between|'
    r'During|Across|Through|Within|Beyond|Above|Below|Plus|Also|And|Or|But)\b', re.I)

# Paraprosdokians: a setup, then a twist ending that makes the reader re-read
# the setup. Hard ban in the register. Two shapes are mechanical enough to find.
#   A. one sentence. Setup, then ", and" / "but" / ";" / ":" / dash, then a
#      short tail that limits or contradicts the setup.
#      "Four steps, and only one of them is yours."
#      "A thousand integrations, and you'll only ever click one."
#   B. two sentences. Any sentence, then a short one that ends on a bare
#      auxiliary or a negation, so it reads as an elided contradiction.
#      "Everything about billing changed. Your invoice didn't."
# Both run on labels as well as prose, because headlines are where these live.
NEG_AUX = (r"isn't|aren't|wasn't|weren't|didn't|doesn't|don't|hasn't|haven't|"
           r"hadn't|won't|wouldn't|can't|couldn't|shouldn't|cannot")
# Words that open a twist tail. Negatives are not here on purpose: "but it
# cannot resell it" is ordinary contrast, and listing them doubled the noise.
LIMITER = re.compile(r"\b(only|just|never|none|nothing|nobody|no one|not one|barely|hardly)\b", re.I)
# Separators, captured so the caller can tell a comma form from a bare one.
TWIST_SEP = re.compile(r'(,?\s+(?:and|but)\s+|,\s+yet\s+|;\s+|:\s+|\s+[-–—]\s+)')  # "yet" needs its comma, else "not yet" trips it
# A tail that ends on a negation, so the verb is elided. "but your invoice didn't."
NEG_END = re.compile(r"\b(" + NEG_AUX + r"|not|never)\s*[.!?]*\s*$", re.I)
# A short second sentence that ends on a bare auxiliary. "It did." "Yours doesn't."
BARE_AUX_END = re.compile(
    r"\b(did|does|do|is|are|was|were|has|have|had|will|would|can|could|should|"
    + NEG_AUX + r")(?:\s+(?:too|either))?\s*[.!?]*\s*$|\b(not|never)\s*[.!?]*\s*$", re.I)
TWIST_OPEN = re.compile(r'^(Not|Only|Just|Never|None|Nobody|No one|Nothing)\b')
URL = re.compile(r'https?://\S+')


def twist_tail(tail, punctuated):
    """Short tail that limits or contradicts the setup. After a comma, colon,
    semicolon or dash the limiter may sit anywhere in the first four words
    ("and you'll only ever click one"). After a bare "and" or "but" it must be
    the first word, or ordinary prose floods the list."""
    words = tail.split()
    if not 1 <= len(words) <= 8:
        return False
    if NEG_END.search(tail):
        return True
    window = ' '.join(words[:4]) if punctuated else words[0]
    return bool(LIMITER.search(window))


def paraprosdokian_a(s):
    """Shape A, one sentence."""
    parts = TWIST_SEP.split(URL.sub(' ', s))
    for i in range(1, len(parts) - 1, 2):
        sep, tail = parts[i], parts[i + 1]
        setup = ''.join(parts[:i])
        punctuated = bool(re.match(r'\s*[,;:\-–—]', sep))
        if len(setup.split()) >= 2 and twist_tail(tail, punctuated):
            return True
    return False


def paraprosdokian_b(a, b):
    """Shape B, two sentences."""
    bw = b.split()
    if not (2 <= len(bw) <= 6 and len(a.split()) >= 3):
        return False
    if paraprosdokian_a(b):          # already reported as shape A, do not double up
        return False
    return bool(BARE_AUX_END.search(b) or (TWIST_OPEN.match(b) and len(bw) <= 4))

def has_verb(s):
    """Deliberately generous. A missed fragment costs less than a false alarm.

    No verb lexicon can ever be complete, so any inflected word counts as
    evidence of a verb. That lets real fragments through, but it keeps the
    report short enough to actually read.
    """
    words = [w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'-]*", s)]
    if any(w in VERBS for w in words):
        return True
    return any(len(w) > 3 and (w.endswith('ed') or w.endswith('ing') or w.endswith('s'))
               for w in words)


def has_verb_strict(s):
    """No -s inflection allowed. Used on short trailing clauses, where an -s
    word is far more likely to be a plural noun than a verb."""
    words = [w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'-]*", s)]
    return (any(w in VERBS for w in words)
            or any(len(w) > 3 and (w.endswith('ed') or w.endswith('ing')) for w in words))


def is_fragment(s):
    """High-confidence only: opens with a determiner or preposition, and
    carries nothing that could be a verb."""
    return FRAGMENT_START.match(s.strip()) and not has_verb(s)

# -------------------------------------------------------------------- report

def report(name, chunks, verbose=False):
    chunks = drop_repeats(chunks)
    prose = [t for k, t in chunks if k == 'prose']
    labels = [t for k, t in chunks if k == 'label']
    everything = ' '.join(t for _, t in chunks)

    sents = [s for block in prose for s in split_sentences(block)]
    real = [s for s in sents if len(s.split()) >= 3]
    words = re.findall(r"[A-Za-z][A-Za-z'-]*", ' '.join(prose))

    print(f'== {name} ==')
    if not real or not words:
        print('  no prose found. Everything read as labels, so only term checks ran.\n')
        wps = spw = 0
    else:
        sy = sum(syl(w) for w in words)
        wps, spw = len(words) / len(real), sy / len(words)
        fk = 0.39 * wps + 11.8 * spw - 15.59
        ease = 206.835 - 1.015 * wps - 84.6 * spw
        print(f'  prose: {len(words)} words in {len(real)} sentences, {wps:.1f} per sentence')
        print(f'  labels: {len(labels)} (headings, captions, footers, table cells: exempt)')
        print(f'  FLOOR ONLY  FK grade {fk:.1f} | reading ease {ease:.1f}'
              f'{"   <-- OVER 12, definitely needs work" if fk > 12 else ""}')
        print('  A good score proves nothing. The checks below are the real test.\n')

    cap = 10 ** 6 if verbose else 40

    def block(title, items, note=''):
        print(f'  -- {title} ({len(items)}) --' + (f'  {note}' if note else ''))
        if not items:
            print('     none')
        for x in items[:cap]:
            print(f'     {x}')
        if len(items) > cap:
            print(f'     ... {len(items) - cap} more (--verbose for all)')
        print()

    block('sentences over 25 words',
          [f'{len(s.split()):>3}w  {s[:110]}' for s in real if len(s.split()) > 25],
          'split or justify each')

    block('colon or dash joining two ideas',
          [s[:110] for s in real
           if re.search(r'\w\s*[:;]\s+\w', s) or re.search(r'\w\s+[-–—]\s+\w', s)],
          'most should become full stops')

    pc = []
    for a, b in zip(real, real[1:]):
        aw, bw = a.split(), b.split()
        if 3 <= len(aw) <= 12 and 3 <= len(bw) <= 12 and aw[0].lower() != bw[0].lower():
            sa = {w.lower().strip('.,') for w in aw}
            sb = {w.lower().strip('.,') for w in bw}
            if len(sa & sb) >= 2 and has_verb(a) and has_verb(b):
                pc.append(f'{a[:55]} || {b[:55]}')
    block('parallel-contrast candidates', pc,
          'keep at most one, only if both halves are plain')

    # PARAPROSDOKIANS, prose and labels ---------------------------------------
    pp = []
    for s in real:
        if paraprosdokian_a(s):
            pp.append(f'A  {s[:110]}')
    for a, b in zip(sents, sents[1:]):
        if paraprosdokian_b(a, b):
            pp.append(f'B  {a[:55]} || {b[:55]}')
    for t in labels:
        ls = split_sentences(t)
        for s in ls:
            if len(s.split()) >= 4 and paraprosdokian_a(s):
                pp.append(f'A  {s[:100]}   (label)')
        for a, b in zip(ls, ls[1:]):
            if paraprosdokian_b(a, b):
                pp.append(f'B  {a[:50]} || {b[:50]}   (label)')
    block('paraprosdokian candidates', pp,
          'hard ban. Setup then twist. Say the setup, then the payoff, as plain sentences')

    # TERM INVENTORY, over labels as well as prose ---------------------------
    first = {}
    for i, s in enumerate(real):
        for w in re.findall(r"[A-Za-z][A-Za-z'-]*", s):
            first.setdefault(w.lower(), i + 1)

    # An all-caps token inside all-caps styling is a design choice, not an
    # acronym. Only count tokens that sit in mixed-case text.
    acro = Counter()
    for _, t in chunks:
        letters = re.sub(r'[^A-Za-z]', '', t)
        if letters and letters.isupper():
            continue
        for w in re.findall(r'\b[A-Z]{2,6}\b', t):
            if w not in ('I', 'A', 'OK'):
                acro[w] += 1
    block('acronyms', [f'{w} x{n}' + (f'  (first in sentence {first[w.lower()]})'
                                      if w.lower() in first else '')
                       for w, n in sorted(acro.items(), key=lambda x: (-x[1], x[0]))],
          'expand each on first use')

    counts = Counter(w.lower() for w in re.findall(r"[A-Za-z][A-Za-z'-]*", everything)
                     if syl(w) >= 3 and w.lower() not in COMMON_LONG and len(w) > 6)
    block('long-word candidates',
          [f'{w} x{n}' + (f'  (first in sentence {first[w]})' if w in first else '')
           for w, n in sorted(counts.items(), key=lambda x: (-x[1], x[0]))],
          'define on first use, or replace with a common word')

    caps = Counter()
    for s in real:
        for w in re.findall(r'(?<=[a-z,)]\s)([A-Z][a-zA-Z]{2,})', s):
            if w.upper() != w:
                caps[w] += 1
    block('capitalised mid-sentence',
          [f'{w} x{n}' for w, n in sorted(caps.items(), key=lambda x: (-x[1], x[0]))],
          'product, tab or defined-term names. Say what each one is')

    block('terms of art',
          [p for p in (x.strip() for x in JARGON_PHRASES) if p and
           re.search(rf'\b{re.escape(p)}\b', everything, re.I)],
          'ordinary words, specialist meaning. Define each')

    block('metaphor-noun candidates',
          [m.replace('_', ' ') for m in METAPHOR_NOUNS
           if re.search(rf'\b{m.replace("_", "[ -]")}\b', everything, re.I)],
          'name the real thing, or define the metaphor')

    # REFERENT CHECK, prose only --------------------------------------------
    refs = []
    for i, s in enumerate(real):
        m = LEADING_REF.match(s)
        if m and len(s.split()) >= 4:
            prev = real[i - 1] if i else '(start of document)'
            refs.append(f'"{m.group(1)}" in: {s[:70]}\n            after: {prev[:70]}')
        for m in SAME_X.finditer(s):
            refs.append(f'"{m.group(0)}" in: {s[:70]}')
    block('reference candidates', refs,
          'referent must be named in this sentence or the one before')

    # VERB CHECK, prose only -------------------------------------------------
    block('fragments (no verb found)',
          [s[:110] for s in real if is_fragment(s)],
          'high-confidence only. Real fragments starting with a short verb are missed')

    tail = []
    for s in real:
        for part in re.split(r',\s+(?:and|but|or)\s+', s)[1:]:
            if len(part.split()) >= 4 and not has_verb_strict(part):
                tail.append(f'{part[:70]}   <- in: {s[:60]}')
    block('clauses with an elided verb', tail, 'write the verb out')

    print('  Judge every candidate above. Most lists will contain false positives.')
    print('  A term is fine if it is defined where it first appears.\n')


def main(argv):
    verbose = '--verbose' in argv
    if '--stdin' in argv:
        report('(stdin)', list(from_text('/dev/stdin')), verbose)
        return 0
    paths = [a for a in argv if not a.startswith('--')]
    if not paths:
        print(__doc__)
        return 1
    for p in paths:
        if not os.path.exists(p):
            print(f'== {p} ==\n  not found\n')
            continue
        report(os.path.basename(p), extract(p), verbose)
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
